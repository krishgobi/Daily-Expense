"""
Excel Report Generator
"""

import io
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, numbers
from openpyxl.chart import PieChart, BarChart, Reference
from openpyxl.chart.series import DataPoint
from openpyxl.utils import get_column_letter


# ── Style helpers ──────────────────────────────────────────────────────────────
def _fill(hex_color: str) -> PatternFill:
    return PatternFill(start_color=hex_color, end_color=hex_color, fill_type="solid")

def _font(bold=False, color="000000", size=10) -> Font:
    return Font(bold=bold, color=color, size=size)

def _border() -> Border:
    s = Side(style="thin", color="E5E7EB")
    return Border(left=s, right=s, top=s, bottom=s)

def _center() -> Alignment:
    return Alignment(horizontal="center", vertical="center", wrap_text=True)

def _left() -> Alignment:
    return Alignment(horizontal="left", vertical="center", wrap_text=True)

HEADER_FILL   = _fill("1F2937")
HEADER_FONT   = _font(bold=True, color="FFFFFF", size=10)
ALT_FILL      = _fill("F9FAFB")
BRAND_FILL    = _fill("3B82F6")
TITLE_FONT    = _font(bold=True, color="1F2937", size=13)
LABEL_FONT    = _font(bold=True, color="374151", size=10)
CURRENCY_FMT  = '#,##0.00'

PIE_COLORS = [
    "3B82F6", "10B981", "F59E0B", "EF4444", "8B5CF6",
    "EC4899", "14B8A6", "F97316", "6366F1", "84CC16",
]


def _apply_header_row(ws, row: int, headers: list[str]):
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=row, column=col, value=h)
        c.fill   = HEADER_FILL
        c.font   = HEADER_FONT
        c.alignment = _center()
        c.border = _border()


def _apply_data_row(ws, row: int, values: list, alt: bool = False):
    fill = ALT_FILL if alt else _fill("FFFFFF")
    for col, v in enumerate(values, 1):
        c = ws.cell(row=row, column=col, value=v)
        c.fill   = fill
        c.border = _border()
        c.alignment = _left()
        c.font   = _font(size=9)


class ExcelReportGenerator:

    @staticmethod
    def generate(report_data: dict, user_name: str) -> bytes:
        wb = Workbook()
        wb.remove(wb.active)

        # ── Sheet 1: Summary ──────────────────────────────────────────────────
        ws = wb.create_sheet("Summary")
        ws.column_dimensions["A"].width = 22
        ws.column_dimensions["B"].width = 26

        ws["A1"] = "Expense Report"
        ws["A1"].font = TITLE_FONT
        ws.merge_cells("A1:B1")

        info = [
            ("Period",    f"{report_data['period_start'].strftime('%d %b %Y')} – {report_data['period_end'].strftime('%d %b %Y')}"),
            ("For",       user_name),
            ("Generated", date.today().strftime("%d %b %Y")),
        ]
        for r, (label, val) in enumerate(info, 3):
            ws.cell(row=r, column=1, value=label).font = LABEL_FONT
            ws.cell(row=r, column=2, value=val).font   = _font(size=10)

        ws["A7"] = "Metric"
        ws["B7"] = "Amount"
        ws["A7"].fill = HEADER_FILL;  ws["A7"].font = HEADER_FONT;  ws["A7"].alignment = _center()
        ws["B7"].fill = HEADER_FILL;  ws["B7"].font = HEADER_FONT;  ws["B7"].alignment = _center()

        totals = [
            ("Total Expenses", report_data["total_expenses"]),
            ("Total Borrowed", report_data["total_borrowed"]),
            ("Total Lent",     report_data["total_lent"]),
        ]
        for i, (label, val) in enumerate(totals, 8):
            ws.cell(row=i, column=1, value=label).font = _font(size=10)
            c = ws.cell(row=i, column=2, value=val)
            c.number_format = CURRENCY_FMT
            c.font = _font(bold=True, size=10)
            if i % 2 == 0:
                ws.cell(row=i, column=1).fill = ALT_FILL
                c.fill = ALT_FILL

        # ── Sheet 2: All Expenses ────────────────────────────────────────────
        we = wb.create_sheet("All Expenses")
        col_widths = [13, 28, 18, 10, 14, 14, 30]
        col_letters = [get_column_letter(i+1) for i in range(7)]
        for cl, w in zip(col_letters, col_widths):
            we.column_dimensions[cl].width = w

        we["A1"] = "All Expenses — Date Wise"
        we["A1"].font = TITLE_FONT
        we.merge_cells("A1:G1")
        we.row_dimensions[1].height = 22

        headers = ["Date", "Purpose", "Category", "Type", "Payment Method", "Amount", "Description"]
        _apply_header_row(we, 3, headers)
        we.row_dimensions[3].height = 18

        expense_rows = report_data.get("expense_rows", [])
        for i, e in enumerate(expense_rows):
            r = i + 4
            alt = i % 2 == 1
            _apply_data_row(we, r, [
                e["date"],
                e["purpose"],
                e["category"],
                e["type"].capitalize(),
                e["payment_method"] or "—",
                e["amount"],
                e["description"] or "",
            ], alt=alt)
            # Format amount column as currency
            c = we.cell(row=r, column=6)
            c.number_format = CURRENCY_FMT
            c.alignment = Alignment(horizontal="right", vertical="center")

        # Total row
        total_row = len(expense_rows) + 4
        we.cell(row=total_row, column=1, value="TOTAL").font = _font(bold=True, size=10)
        tc = we.cell(row=total_row, column=6, value=report_data["total_expenses"])
        tc.number_format = CURRENCY_FMT
        tc.font = _font(bold=True, size=10)
        tc.fill = _fill("DBEAFE")

        # Freeze header
        we.freeze_panes = "A4"

        # ── Sheet 3: Category Breakdown + Pie Chart ───────────────────────────
        wc = wb.create_sheet("Categories")
        wc.column_dimensions["A"].width = 22
        wc.column_dimensions["B"].width = 16
        wc.column_dimensions["C"].width = 14
        wc.column_dimensions["D"].width = 10

        wc["A1"] = "Spending by Category"
        wc["A1"].font = TITLE_FONT
        wc.merge_cells("A1:D1")

        _apply_header_row(wc, 3, ["Category", "Amount", "% of Total", "Expenses"])

        breakdown = report_data.get("category_breakdown", [])
        for i, cat in enumerate(breakdown):
            r = i + 4
            alt = i % 2 == 1
            _apply_data_row(wc, r, [
                cat["name"],
                cat["amount"],
                cat["percentage"] / 100,
                cat["count"],
            ], alt=alt)
            wc.cell(row=r, column=2).number_format = CURRENCY_FMT
            wc.cell(row=r, column=3).number_format = "0.0%"

        # Pie chart
        if breakdown:
            pie = PieChart()
            pie.title = "Category Breakdown"
            pie.style = 10
            pie.width  = 14
            pie.height = 10

            # Data: amounts from column B
            data_ref = Reference(wc, min_col=2, min_row=3, max_row=3 + len(breakdown))
            cats_ref = Reference(wc, min_col=1, min_row=4, max_row=3 + len(breakdown))
            pie.add_data(data_ref, titles_from_data=True)
            pie.set_categories(cats_ref)
            pie.dataLabels = None

            # Colour each slice
            series = pie.series[0]
            for idx, cat in enumerate(breakdown):
                pt = DataPoint(idx=idx)
                pt.graphicalProperties.solidFill = PIE_COLORS[idx % len(PIE_COLORS)]
                series.dPt.append(pt)

            wc.add_chart(pie, "F3")

        # ── Save ──────────────────────────────────────────────────────────────
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf.getvalue()
